#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
화장품 성분표 웹 도구
사용법: python server.py
→ 브라우저가 자동으로 열립니다.
"""

import io, json, sys, threading, time, webbrowser
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse

BASE_DIR = Path(__file__).parent
sys.path.insert(0, str(BASE_DIR))

try:
    import cosmetic_tool as ct
except ImportError:
    sys.exit("cosmetic_tool.py 가 같은 폴더에 있어야 합니다.")

try:
    import requests as _req
except ImportError:
    sys.exit("pip install requests 를 먼저 실행하세요.")

CACHE_PATH = BASE_DIR / "cache" / "ingredients.json"

_state = {"running": False, "progress": "", "error": "", "done": False}
_payload = None
_index: dict = {}
_field_map: dict = {}


def _load_cache():
    global _payload, _index, _field_map
    if CACHE_PATH.exists():
        try:
            _payload = json.loads(CACHE_PATH.read_text(encoding="utf-8"))
            _index, _field_map = ct.build_index(_payload)
            return True
        except Exception:
            pass
    return False


def _fetch_thread():
    global _payload, _index, _field_map
    _state.update({"running": True, "error": "", "done": False, "progress": "연결 중..."})
    try:
        session = _req.Session()
        session.headers["User-Agent"] = "cosmetic-tool/2.0"
        key = ct.load_service_key()

        _state["progress"] = "오퍼레이션 탐색 중..."
        op = ct.detect_operation(session, key)
        if not op:
            _state["error"] = "오퍼레이션을 찾을 수 없습니다. probe 를 실행해 확인하세요."
            return

        all_items, page, fm = [], 1, {}
        while True:
            code, msg, items, total = ct.call_api(session, key, op, page, 100)
            if not items:
                if code not in ("00", ""):
                    _state["error"] = f"API 오류: {code} — {msg}"
                break
            if not fm:
                fm = ct.detect_fields(items)
            all_items.extend(items)
            _state["progress"] = f"{len(all_items):,} / {total:,} 건"
            if total and len(all_items) >= total:
                break
            page += 1
            time.sleep(0.2)

        if not all_items:
            if not _state["error"]:
                _state["error"] = "데이터를 받지 못했습니다."
            return

        _payload = {"operation": op, "field_map": fm, "count": len(all_items), "items": all_items}
        CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
        CACHE_PATH.write_text(json.dumps(_payload, ensure_ascii=False), encoding="utf-8")
        _index, _field_map = ct.build_index(_payload)
        _state.update({"done": True, "progress": f"완료: {len(all_items):,}건 로드됨"})

    except SystemExit as e:
        _state["error"] = str(e)
    except Exception as e:
        _state["error"] = f"{type(e).__name__}: {e}"
    finally:
        _state["running"] = False


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        pass

    def _json(self, data, status=200):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _file(self, data: bytes, mime: str, name: str):
        self.send_response(200)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Disposition", f'attachment; filename*=UTF-8\'\'{name}')
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        path = urlparse(self.path).path

        if path in ("/", "/index.html"):
            try:
                html = (BASE_DIR / "index.html").read_bytes()
            except FileNotFoundError:
                self._json({"error": "index.html 없음"}, 500)
                return
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(html)))
            self.end_headers()
            self.wfile.write(html)

        elif path == "/api/status":
            count = 0
            if _payload:
                count = _payload.get("count", 0)
            elif CACHE_PATH.exists():
                try:
                    count = json.loads(CACHE_PATH.read_text("utf-8")).get("count", 0)
                except Exception:
                    pass
            self._json({
                "cached": CACHE_PATH.exists(),
                "count": count,
                "running": _state["running"],
                "progress": _state["progress"],
                "error": _state["error"],
                "done": _state["done"],
            })

        elif path == "/api/fetch":
            if _state["running"]:
                self._json({"ok": False, "msg": "이미 진행 중입니다"})
                return
            _state["done"] = False
            threading.Thread(target=_fetch_thread, daemon=True).start()
            self._json({"ok": True})

        else:
            self.send_response(404)
            self.end_headers()

    def do_POST(self):
        length = int(self.headers.get("Content-Length", 0))
        body = self.rfile.read(length)
        path = urlparse(self.path).path

        if path == "/api/lookup":
            if not _index:
                _load_cache()
            if not _index:
                self._json({"error": "DB가 없습니다. 먼저 DB를 초기화하세요."}, 400)
                return
            try:
                names = [n.strip() for n in json.loads(body).get("names", []) if str(n).strip()]
                results = []
                for name in names:
                    info, status = ct.lookup_one(name, _index, _field_map)
                    results.append({
                        "name": name,
                        "eng": (info or {}).get("eng", ""),
                        "cas": (info or {}).get("cas", ""),
                        "defn": (info or {}).get("defn", ""),
                        "status": status,
                    })
                self._json({"results": results})
            except Exception as e:
                self._json({"error": str(e)}, 500)

        elif path == "/api/export":
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                rows = json.loads(body).get("rows", [])
                wb = Workbook()
                ws = wb.active
                ws.title = "성분표"

                thin = Side(style="thin", color="D1D5DB")
                bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
                h_fill = PatternFill("solid", fgColor="7C3AED")

                headers = ["한글성분명", "영문명", "성분비율(%)", "CAS.NO", "기능", "매칭상태"]
                ws.append(headers)
                for i in range(1, len(headers) + 1):
                    c = ws.cell(1, i)
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = h_fill
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.border = bdr
                ws.row_dimensions[1].height = 22

                fills = {
                    "찾음": "D1FAE5", "유사일치": "FEF3C7",
                    "못찾음": "FEE2E2", "다중일치(첫번째)": "DBEAFE",
                }
                for ri, r in enumerate(rows, 2):
                    ws.append([r.get("name",""), r.get("eng",""), r.get("ratio",""),
                                r.get("cas",""), r.get("defn",""), r.get("status","")])
                    for ci in range(1, 7):
                        cell = ws.cell(ri, ci)
                        cell.border = bdr
                        cell.alignment = Alignment(vertical="center", wrap_text=(ci == 5))
                    st = r.get("status","")
                    fg = fills.get(st)
                    if fg:
                        ws.cell(ri, 6).fill = PatternFill("solid", fgColor=fg)

                for i, w in enumerate([22, 30, 11, 16, 45, 14], 1):
                    ws.column_dimensions[chr(64 + i)].width = w

                buf = io.BytesIO()
                wb.save(buf)
                from urllib.parse import quote
                self._file(buf.getvalue(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           quote("성분표.xlsx"))
            except Exception as e:
                self._json({"error": str(e)}, 500)

        else:
            self.send_response(404)
            self.end_headers()


if __name__ == "__main__":
    _load_cache()
    PORT = 8282
    httpd = ThreadingHTTPServer(("localhost", PORT), Handler)
    url = f"http://localhost:{PORT}"
    print(f"\n{'─'*44}")
    print(f"  화장품 성분표 도구  →  {url}")
    print(f"  종료: Ctrl+C")
    print(f"{'─'*44}\n")
    threading.Timer(1.0, lambda: webbrowser.open(url)).start()
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\n서버 종료.")
